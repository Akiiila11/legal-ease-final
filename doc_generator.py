import os
import re

from typing import Optional
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from fpdf import FPDF
    _FPDF_AVAILABLE = True
except ImportError:
    _FPDF_AVAILABLE = False


# --------------------------------------------------------------------------
# CONFIGURATION
# --------------------------------------------------------------------------

OUTPUT_DIR = "generated_docs"
os.makedirs(OUTPUT_DIR, exist_ok=True)


TITLES = {
    "employment_contract": "EMPLOYMENT CONTRACT",
    "nda": "NON-DISCLOSURE AGREEMENT",
    "lease_agreement": "LEASE AGREEMENT",
}


# --------------------------------------------------------------------------
# COMMON HELPERS
# --------------------------------------------------------------------------

def _key_terms(doc_type: str, data: dict) -> dict:
    """Return important input values for the generated document."""

    if doc_type == "employment_contract":
        return {
            "Company": data.get("company_name", ""),
            "Employee": data.get("employee_name", ""),
            "Job Title": data.get("job_title", ""),
            "Start Date": data.get("start_date", ""),
            "Compensation": data.get("compensation", ""),
            "Work Location": data.get("work_location", ""),
            "Employment Type": data.get("employment_type", ""),
        }

    if doc_type == "nda":
        return {
            "Disclosing Party": data.get("disclosing_party", ""),
            "Receiving Party": data.get("receiving_party", ""),
            "Purpose": data.get("purpose", ""),
            "Scope": data.get("scope", ""),
            "Effective Date": data.get("effective_date", ""),
            "Term": data.get("term", ""),
        }

    if doc_type == "lease_agreement":
        return {
            "Landlord": data.get("landlord_name", ""),
            "Tenant": data.get("tenant_name", ""),
            "Property Address": data.get("property_address", ""),
            "Lease Start": data.get("lease_start", ""),
            "Lease End": data.get("lease_end", ""),
            "Monthly Rent": data.get("monthly_rent", ""),
            "Security Deposit": data.get("security_deposit", ""),
        }

    return {}


def _split_sections(raw_text: str) -> list[tuple[str, str]]:
    """
    Split generated legal text into heading/body sections.

    Expected format:

    1. Parties and Effective Date
    Text...

    2. Position and Duties
    Text...
    """

    if not raw_text:
        return []

    lines = raw_text.splitlines()

    sections = []
    current_heading = None
    current_body = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            if current_heading is not None:
                current_body.append("")
            continue

        # Detect numbered headings such as:
        # 1. Parties and Effective Date
        # 2. Position and Duties
        if re.match(r"^\d+\.\s+.+", stripped):
            if current_heading is not None:
                sections.append(
                    (
                        current_heading,
                        "\n".join(current_body).strip()
                    )
                )

            current_heading = stripped
            current_body = []

        else:
            if current_heading is None:
                current_heading = "Agreement"
                current_body = []

            current_body.append(stripped)

    if current_heading is not None:
        sections.append(
            (
                current_heading,
                "\n".join(current_body).strip()
            )
        )

    return sections


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def build_docx(
    doc_type: str,
    data: dict,
    raw_text: str,
    filename: str,
    logo_path: Optional[str] = None
) -> str:

    document = Document()

    # --------------------------------------------------
    # Page setup
    # --------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # --------------------------------------------------
    # Logo
    # --------------------------------------------------

    if logo_path and os.path.exists(logo_path):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(
            logo_path,
            width=Inches(1.5)
        )

    # --------------------------------------------------
    # Title
    # --------------------------------------------------

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        TITLES.get(doc_type, "LEGAL DOCUMENT")
    )

    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph()

    # --------------------------------------------------
    # Key Terms
    # --------------------------------------------------

    terms = _key_terms(doc_type, data)

    if terms:
        heading = document.add_paragraph()

        run = heading.add_run("Key Terms")
        run.bold = True
        run.font.size = Pt(12)

        for key, value in terms.items():
            paragraph = document.add_paragraph()

            label = paragraph.add_run(f"{key}: ")
            label.bold = True

            paragraph.add_run(str(value))

        document.add_paragraph()

    # --------------------------------------------------
    # Agreement Text
    # --------------------------------------------------

    sections = _split_sections(raw_text)

    for heading_text, body in sections:

        # Section heading
        heading = document.add_paragraph()

        run = heading.add_run(heading_text)
        run.bold = True
        run.font.size = Pt(11)

        # Section body
        if body:
            for paragraph_text in body.split("\n"):

                paragraph_text = paragraph_text.strip()

                if not paragraph_text:
                    continue

                paragraph = document.add_paragraph(
                    paragraph_text
                )

                paragraph.paragraph_format.space_after = Pt(6)

                for run in paragraph.runs:
                    run.font.size = Pt(10)

    # --------------------------------------------------
    # Save
    # --------------------------------------------------

    out_path = os.path.join(
        OUTPUT_DIR,
        f"{filename}.docx"
    )

    document.save(out_path)

    return out_path


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

if _FPDF_AVAILABLE:

    class LegalPDF(FPDF):

        def header(self):
            self.set_font(
                "Helvetica",
                "B",
                14
            )

            self.cell(
                0,
                10,
                self.title_text,
                align="C",
                new_x="LMARGIN",
                new_y="NEXT"
            )

            self.ln(2)

        def footer(self):
            self.set_y(-15)

            self.set_font(
                "Helvetica",
                "I",
                8
            )

            self.cell(
                0,
                10,
                f"Page {self.page_no()}",
                align="C"
            )


def _sanitize_pdf_text(text: str) -> str:
    """
    Clean AI-generated text before sending it to FPDF.
    """

    if not text:
        return ""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u00a0": " ",
        "\u2022": "-",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Remove unsupported control characters
    text = "".join(
        char
        if char in "\n\r\t" or ord(char) >= 32
        else " "
        for char in text
    )

    # Break extremely long words/URLs
    text = re.sub(
        r"(\S{40})(?=\S)",
        r"\1 ",
        text
    )

    return text.strip()


def build_pdf(
    doc_type: str,
    data: dict,
    raw_text: str,
    filename: str
) -> str:

    if not _FPDF_AVAILABLE:
        raise RuntimeError(
            "fpdf2 is not installed. Run: pip install fpdf2"
        )

    pdf = LegalPDF()

    pdf.title_text = TITLES.get(
        doc_type,
        "LEGAL DOCUMENT"
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )

    pdf.add_page()

    # --------------------------------------------------
    # Key Terms
    # --------------------------------------------------

    terms = _key_terms(
        doc_type,
        data
    )

    if terms:

        pdf.set_font(
            "Helvetica",
            "B",
            12
        )

        pdf.cell(
            0,
            8,
            "Key Terms",
            new_x="LMARGIN",
            new_y="NEXT"
        )

        pdf.set_font(
            "Helvetica",
            "",
            10
        )

        for key, value in terms.items():

            value = _sanitize_pdf_text(
                str(value)
            )

            pdf.multi_cell(
                0,
                6,
                f"{key}: {value}",
                new_x="LMARGIN",
                new_y="NEXT"
            )

        pdf.ln(4)

    # --------------------------------------------------
    # Agreement Text
    # --------------------------------------------------

    for heading, body in _split_sections(raw_text):

        # Heading
        pdf.set_font(
            "Helvetica",
            "B",
            11
        )

        heading = _sanitize_pdf_text(
            heading
        )

        pdf.multi_cell(
            0,
            7,
            heading,
            new_x="LMARGIN",
            new_y="NEXT"
        )

        # Body
        pdf.set_font(
            "Helvetica",
            "",
            10
        )

        body = _sanitize_pdf_text(
            body
        )

        for paragraph in body.split("\n"):

            paragraph = paragraph.strip()

            if not paragraph:
                continue

            pdf.multi_cell(
                0,
                6,
                paragraph,
                new_x="LMARGIN",
                new_y="NEXT"
            )

            pdf.ln(1)

        pdf.ln(2)

    # --------------------------------------------------
    # Save
    # --------------------------------------------------

    out_path = os.path.join(
        OUTPUT_DIR,
        f"{filename}.pdf"
    )

    pdf.output(out_path)

    return out_path


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------

def build_txt(
    doc_type: str,
    raw_text: str,
    filename: str
) -> str:

    out_path = os.path.join(
        OUTPUT_DIR,
        f"{filename}.txt"
    )

    with open(
        out_path,
        "w",
        encoding="utf-8"
    ) as f:

        f.write(
            f"{TITLES.get(doc_type, 'LEGAL DOCUMENT')}\n"
        )

        f.write(
            "=" * 60 + "\n\n"
        )

        f.write(raw_text)

    return out_path


# --------------------------------------------------------------------------
# GENERATE ALL FORMATS
# --------------------------------------------------------------------------

def generate_all_formats(
    doc_type: str,
    data: dict,
    raw_text: str,
    filename: str,
    logo_path: Optional[str] = None
) -> dict:

    outputs = {
        "docx": build_docx(
            doc_type,
            data,
            raw_text,
            filename,
            logo_path
        ),

        "txt": build_txt(
            doc_type,
            raw_text,
            filename
        ),
    }

    if _FPDF_AVAILABLE:
        outputs["pdf"] = build_pdf(
            doc_type,
            data,
            raw_text,
            filename
        )

    return outputs